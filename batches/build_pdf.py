"""Render the lead list, contacts and outreach messages into Gliped_Lead_Pack.pdf.

Reads the Leads tab of Gliped_Outreach_Tracker.xlsx, so run build_tracker.py first.
Run from batches/:  python3 build_pdf.py
"""
import datetime as dt
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

OUT = "../Gliped_Lead_Pack.pdf"
pdfmetrics.registerFont(TTFont("DV", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DVB", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))

NAVY = colors.HexColor("#1F3A5F")
GREY = colors.HexColor("#667085")
LINE = colors.HexColor("#D0D5DD")
GREEN = colors.HexColor("#D1FADF")
RED = colors.HexColor("#F8D7DA")
AMBER = colors.HexColor("#FFF6CC")

S = {
    "title": ParagraphStyle("t", fontName="DVB", fontSize=20, leading=24, textColor=NAVY),
    "sub": ParagraphStyle("s", fontName="DV", fontSize=10, leading=14, textColor=GREY),
    "h1": ParagraphStyle("h1", fontName="DVB", fontSize=13, leading=16, textColor=NAVY, spaceBefore=6, spaceAfter=4),
    "h2": ParagraphStyle("h2", fontName="DVB", fontSize=11, leading=14, textColor=NAVY),
    "body": ParagraphStyle("b", fontName="DV", fontSize=8.5, leading=11.5, alignment=TA_LEFT),
    "small": ParagraphStyle("sm", fontName="DV", fontSize=7.5, leading=9.5),
    "cell": ParagraphStyle("c", fontName="DV", fontSize=7, leading=8.6),
    "cellb": ParagraphStyle("cb", fontName="DVB", fontSize=7, leading=8.6, textColor=colors.white),
    "label": ParagraphStyle("l", fontName="DVB", fontSize=8, leading=10, textColor=NAVY, spaceAfter=2),
    "quote": ParagraphStyle("q", fontName="DV", fontSize=8.5, leading=11.5, leftIndent=6,
                            borderColor=LINE, borderWidth=0, backColor=colors.HexColor("#F5F7FA"), borderPadding=4,
                            spaceBefore=6, spaceAfter=6),
}


def P(text, style="body"):
    return Paragraph(escape(str(text or "")).replace("\n", "<br/>"), S[style])


def load():
    ws = load_workbook("../Gliped_Outreach_Tracker.xlsx")["Leads"]
    head = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        d = dict(zip(head, r))
        s = d["Score"] or 0
        d["Tier"] = "Priority" if s >= 60 else ("Warm" if s >= 40 else "Nurture")
        for k in ("Angle", "Teardown focus", "Activity notes", "Notes", "Trigger"):
            v = d.get(k)
            if isinstance(v, str) and v:
                d[k] = v[0].upper() + v[1:]
        rows.append(d)
    return rows


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("DV", 7)
    canvas.setFillColor(GREY)
    canvas.drawString(12 * mm, 7 * mm, "Gliped lead pack. Prepared for Prabal. Contains personal contact data, share carefully.")
    canvas.drawRightString(doc.pagesize[0] - 12 * mm, 7 * mm, f"Page {doc.page}")
    canvas.restoreState()


leads = load()
found = sum(1 for l in leads if l["Email status"] == "Verified")
notfound = sum(1 for l in leads if l["Email status"] == "Not found")
pending = sum(1 for l in leads if l["Email status"] == "Not checked")

doc = SimpleDocTemplate(OUT, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                        topMargin=12 * mm, bottomMargin=13 * mm, title="Gliped lead pack",
                        author="Gliped")
W = doc.width
story = []

# ---------- cover / summary ----------
story += [P("Gliped lead pack", "title"),
          P(f"{len(leads)} founders and CEOs across India, the UK, Canada and the US. Batches 01 and 02. "
            f"Generated {dt.date.today():%d %b %Y}.", "sub"),
          Spacer(1, 8)]

def count(key, val):
    return sum(1 for l in leads if l[key] == val)

stats = [["Leads", "Priority (60+)", "Warm (40 to 59)", "Nurture", "Verified emails", "Email not found", "Email pending"],
         [len(leads), count("Tier", "Priority"), count("Tier", "Warm"), count("Tier", "Nurture"), found, notfound, pending]]
t = Table(stats, colWidths=[W / 7] * 7)
t.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "DVB", 8), ("FONT", (0, 1), (-1, 1), "DVB", 16),
                       ("TEXTCOLOR", (0, 0), (-1, 0), GREY), ("TEXTCOLOR", (0, 1), (-1, 1), NAVY),
                       ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("BOX", (0, 0), (-1, -1), 0.5, LINE),
                       ("INNERGRID", (0, 0), (-1, -1), 0.5, LINE), ("TOPPADDING", (0, 0), (-1, -1), 6),
                       ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
story += [t, Spacer(1, 10)]

region_rows = [["Region", "Leads", "Verified emails", "Priority"]]
for reg in ["India", "UK", "Canada", "US"]:
    rl = [l for l in leads if l["Region"] == reg]
    region_rows.append([reg, len(rl), sum(1 for l in rl if l["Email status"] == "Verified"),
                        sum(1 for l in rl if l["Tier"] == "Priority")])
t = Table(region_rows, colWidths=[40 * mm, 30 * mm, 35 * mm, 30 * mm])
t.setStyle(TableStyle([("FONT", (0, 0), (-1, 0), "DVB", 8), ("FONT", (0, 1), (-1, -1), "DV", 8.5),
                       ("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                       ("GRID", (0, 0), (-1, -1), 0.5, LINE), ("ALIGN", (1, 0), (-1, -1), "CENTER")]))
story += [t, Spacer(1, 10)]

story += [P("How to use this pack", "h1"),
          P("1. Start with Priority leads. Warm up on LinkedIn for a few days, then send the connection note, then the first DM once they accept.\n"
            "2. If the connection isn't accepted within 5 days, or there's no reply to the first DM, send the cold email to the verified address. One email follow-up after 5 days, on the same thread.\n"
            "3. Before sending, spend 30 seconds on their recent LinkedIn activity. Where a message says [one line on ...], read that post and write the line yourself.\n"
            "4. \"Verify profile\" means the LinkedIn link is a people search. Match the company and role first.\n"
            "5. Log every step in the Tracker tab of Gliped_Outreach_Tracker.xlsx. This PDF is a reading copy; the workbook is the working file."),
          Spacer(1, 6),
          P("Emails come from Prospeo and were requested as verified only (SMTP or BounceBan checked). Not found means Prospeo had no verified address, "
            "and Pending means Prospeo's rate limit stopped the lookup before it reached them.", "small")]
story.append(PageBreak())

# ---------- contact table ----------
story.append(P("Contact list", "h1"))
hdr = ["#", "Name", "Role, company", "Region", "Tier", "Email", "Email status", "LinkedIn"]
data = [[Paragraph(h, S["cellb"]) for h in hdr]]
style = [("BACKGROUND", (0, 0), (-1, 0), NAVY), ("GRID", (0, 0), (-1, -1), 0.4, LINE),
         ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2.5),
         ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5)]
for i, l in enumerate(leads, 1):
    li = l["LinkedIn URL"] or ""
    li_txt = "people search (verify)" if "search/results" in li else li.replace("https://www.linkedin.com", "")
    data.append([P(i, "cell"), P(l["Name"], "cell"), P(f'{l["Role"]}, {l["Company"]}', "cell"), P(l["Region"], "cell"),
                 P(l["Tier"], "cell"), P(l["Email address"] or "", "cell"), P(l["Email status"], "cell"), P(li_txt, "cell")])
    colour = {"Verified": GREEN, "Not found": RED, "Not checked": AMBER}.get(l["Email status"])
    if colour:
        style.append(("BACKGROUND", (6, i), (6, i), colour))
cw = [8 * mm, 32 * mm, 62 * mm, 16 * mm, 15 * mm, 55 * mm, 20 * mm, W - 208 * mm]
t = Table(data, colWidths=cw, repeatRows=1)
t.setStyle(TableStyle(style))
story += [t, PageBreak()]

# ---------- lead cards ----------
story.append(P("Lead cards", "h1"))
for i, l in enumerate(leads, 1):
    tag = "" if l["Profile confirmed"] == "Yes" else "  (verify profile)"
    head = Table([[P(f'{i}. {l["Name"]}, {l["Role"]}, {l["Company"]}{tag}', "h2"),
                   P(f'{l["Region"]} | {l["Tier"]} | score {l["Score"]} | batch {l["Batch"]}', "small")]],
                 colWidths=[W * 0.7, W * 0.3])
    head.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, 0), 0.8, NAVY), ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                              ("ALIGN", (1, 0), (1, 0), "RIGHT"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    email_line = l["Email address"] or {"Not found": "No verified email found", "Not checked": "Pending lookup"}.get(l["Email status"], "")
    left = [P("Contact", "label"), P(f'Email: {email_line}\nLinkedIn: {l["LinkedIn URL"]}', "small"), Spacer(1, 3),
            P("Trigger", "label"), P(l["Trigger"], "small"), Spacer(1, 3),
            P("LinkedIn activity", "label"), P(f'{l["Activity pattern"]}. {l["Activity notes"] or ""}', "small"), Spacer(1, 3),
            P("Angle", "label"), P(l["Angle"], "small"), Spacer(1, 3),
            P("Teardown focus", "label"), P(l["Teardown focus"], "small")]
    if l["Notes"]:
        left += [Spacer(1, 3), P("Notes", "label"), P(l["Notes"], "small")]
    if l["Email lookup note"]:
        left += [Spacer(1, 3), P("Email note", "label"), P(l["Email lookup note"], "small")]
    right = [P("Connection note", "label"), P(l["Connection note"], "quote"), Spacer(1, 4),
             P("First DM", "label"), P(l["First DM"], "quote"), Spacer(1, 4),
             P(f'Cold email. Subject: {l["Email subject"]}', "label"), P(l["Email body"], "quote")]
    body = Table([[left, right]], colWidths=[W * 0.42, W * 0.58])
    body.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("RIGHTPADDING", (0, 0), (0, 0), 10)]))
    story += [KeepTogether([head, Spacer(1, 4), body]), Spacer(1, 12)]
    if i < len(leads):
        story.append(PageBreak())

doc.build(story, onFirstPage=footer, onLaterPages=footer)
print("saved", OUT, len(leads), "leads", found, "verified", notfound, "not found", pending, "pending")
