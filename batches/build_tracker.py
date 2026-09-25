"""Build Gliped_Outreach_Tracker.xlsx from batch 01 (markdown + CSV) and batch 02 (batch02_data.py).

Run from the batches/ folder:  python3 build_tracker.py
"""
import csv
import re
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import batch02_data as b2

OUT = "../Gliped_Outreach_Tracker.xlsx"
FONT = "Arial"

PATTERN_LABELS = {"milestone-only": "Milestone-only", "went-quiet": "Went quiet", "company-news-voice": "Company-news voice",
                  "emerging-writer": "Emerging writer", "already-consistent": "Already consistent"}

# ---------- batch 01: parse lead cards from the markdown, join with CSV ----------
def parse_batch01():
    text = open("2026-09-22-batch-01.md").read()
    cards = re.split(r"\n### \d+\. ", text.split("## Lead cards", 1)[1])[1:]
    rows = {r["name"]: r for r in csv.DictReader(open("../gliped-crm-processed-tracker.csv"))}
    out = []
    for card in cards:
        card = card.split("\n---")[0].split("\n## ")[0]
        head, body = card.split("\n", 1)
        name = head.split(",")[0].strip()
        fields, cur = {}, None
        for line in body.split("\n"):
            m = re.match(r"- \*\*(.+?):\*\*\s?(.*)", line)
            if m:
                cur = m.group(1)
                fields[cur] = m.group(2).strip()
            elif line.startswith(">") and cur:
                fields[cur] = (fields[cur] + "\n" + line.lstrip("> ").rstrip()).strip()
        r = rows[name]
        notes = " ".join(x for x in [fields.get("Note", ""), fields.get("Region note", ""), r["notes"]] if x)
        out.append(dict(
            name=name, role=r["role"], company=r["company"], region=r["region"], city="",
            url=r["linkedin_url"], confirmed=True,
            trigger=fields.get("Trigger", r["trigger"]), trigger_date=r["trigger"],
            source="", pattern=PATTERN_LABELS[r["activity_pattern"]],
            last_post=r["last_indexed_own_post"],
            activity=fields.get("Activity pattern", ""), angle=fields.get("Angle", ""),
            score=int(r["score"]), connect=fields.get("Connection note", ""),
            dm=fields.get("First DM", ""), teardown=fields.get("Teardown focus", ""),
            notes=notes.strip(), batch="01", date="2026-09-22",
        ))
    assert len(out) == 20, len(out)
    return out


def batch02():
    out = []
    for l in b2.LEADS:
        d = dict(l)
        d["batch"], d["date"] = b2.BATCH, b2.DATE
        out.append(d)
    return out


def trigger_month(l):
    """Batch 01 CSV holds trigger text, not dates; pull 'Mon YYYY' if present."""
    if l["batch"] == "02":
        return l["trigger_date"]
    m = re.search(r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\w* (20\d\d)", l["trigger_date"])
    if m:
        mon = dt.datetime.strptime(m.group(1), "%b").month
        return f"{m.group(2)}-{mon:02d}"
    return ""


leads = parse_batch01() + batch02()
for i, l in enumerate(leads, 1):
    l["id"] = f"G{i:03d}"

# ---------- styles ----------
thin = Side(style="thin", color="D0D5DD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFF6CC")
LINK_FONT = Font(name=FONT, size=10, color="1155CC", underline="single")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def header(ws, cols, row=1):
    for c, (title, width) in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=title)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[row].height = 32


wb = Workbook()

# ---------- How to use ----------
ws = wb.active
ws.title = "How to use"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 100
rows = [
    ("Gliped outreach tracker", ""),
    ("", ""),
    ("Tabs", ""),
    ("Leads", "Every lead with trigger, LinkedIn activity pattern, score and ready-to-send messages. One row per person. Add new batches at the bottom with the next ID (G061, G062, ...)."),
    ("Tracker", "Your working sheet. One row per lead ID. Fill the yellow cells as you go. Names, company, region and tier pull from Leads automatically by ID."),
    ("Dashboard", "Counts and conversion rates by stage, region and activity pattern. Updates itself from Tracker."),
    ("Sequence", "The shared outreach cadence, follow-up messages and send times."),
    ("Lists", "Values behind the drop-downs. Edit here to change the options."),
    ("", ""),
    ("How to fill the Tracker", ""),
    ("Status", "Pick from the drop-down. Move it forward each time something happens."),
    ("Date columns", "Type dates as YYYY-MM-DD (for example 2026-09-29). Follow-up due dates calculate themselves."),
    ("FU1 due", "First DM sent + 4 days. Turns red when overdue and FU1 sent is empty."),
    ("FU2 due", "FU1 sent + 6 days. Turns red when overdue and FU2 sent is empty."),
    ("Days since last touch", "Days since the latest date you've entered in that row."),
    ("Yellow cells", "The only cells you need to edit. Grey or white cells are formulas or reference data."),
    ("", ""),
    ("Example row (format only)", "G999 | Status: DM sent | Warm-up started 2026-09-26 | Connection sent 2026-10-01 | Accepted 2026-10-02 | First DM sent 2026-10-02 | FU1 due fills in as 2026-10-06 | Next action: send FU1 teardown offer"),
    ("", ""),
    ("Scoring", "Trigger in last 90 days (max 30) + profile gap (max 20) + fit (10) + active in last 30 days (10). Intent signals were not visible, so add +30 by hand when someone views your profile or engages with Gliped. Tier: 60+ Priority, 40 to 59 Warm, under 40 Nurture."),
    ("Profile confirmed = No", "The LinkedIn link is a people search. Open it, match company and role, then paste the real profile URL into Leads and set Profile confirmed to Yes."),
    ("Activity data", "Posting timelines come from publicly indexed LinkedIn posts, dated from each post's activity ID. The index only catches some posts, so 'last indexed own post' means at least that recent."),
]
for r, (a, bval) in enumerate(rows, 1):
    ws.cell(row=r, column=1, value=a).font = Font(name=FONT, bold=bool(a) and not bval or r == 1, size=14 if r == 1 else 10)
    c = ws.cell(row=r, column=2, value=bval)
    c.font, c.alignment = BODY_FONT, WRAP
    ws.cell(row=r, column=1).alignment = TOP
    if a in ("Leads", "Tracker", "Dashboard", "Sequence", "Lists", "Status", "Date columns", "FU1 due", "FU2 due",
             "Days since last touch", "Yellow cells", "Example row (format only)", "Scoring", "Profile confirmed = No", "Activity data"):
        ws.cell(row=r, column=1).font = Font(name=FONT, bold=True, size=10)
ws["B16"].fill = INPUT_FILL

# ---------- Lists ----------
wl = wb.create_sheet("Lists")
STATUSES = ["Not started", "Warming up", "Connection sent", "Connected", "DM sent", "Follow-up 1 sent",
            "Follow-up 2 sent", "Replied", "Teardown sent", "Call booked", "Call held", "Proposal sent",
            "Won", "Lost", "No response", "Nurture"]
OUTCOMES = ["", "Won", "Lost", "Not now", "No fit", "Referred"]
YN = ["Yes", "No"]
for col, (title, vals) in enumerate([("Status", STATUSES), ("Outcome", OUTCOMES[1:]), ("Yes/No", YN)], 1):
    wl.cell(row=1, column=col, value=title).font = HEAD_FONT
    wl.cell(row=1, column=col).fill = HEAD_FILL
    wl.column_dimensions[get_column_letter(col)].width = 20
    for r, v in enumerate(vals, 2):
        wl.cell(row=r, column=col, value=v).font = BODY_FONT

# ---------- Leads ----------
wsL = wb.create_sheet("Leads", 1)
LCOLS = [("ID", 7), ("Batch", 7), ("Date added", 11), ("Name", 22), ("Role", 20), ("Company", 24), ("Region", 10),
         ("City", 13), ("LinkedIn URL", 30), ("Profile confirmed", 10), ("Trigger", 42), ("Trigger date", 11),
         ("Source", 24), ("Activity pattern", 17), ("Last indexed own post", 12), ("Activity notes", 50),
         ("Score", 7), ("Tier", 9), ("Angle", 45), ("Connection note", 45), ("First DM", 45),
         ("Teardown focus", 40), ("Notes", 36)]
header(wsL, LCOLS)
for r, l in enumerate(leads, 2):
    vals = [l["id"], l["batch"], l["date"], l["name"], l["role"], l["company"], l["region"], l["city"],
            l["url"], "Yes" if l["confirmed"] else "No", l["trigger"], trigger_month(l), l["source"],
            l["pattern"], l["last_post"], l["activity"], l["score"], None, l["angle"], l["connect"],
            l["dm"], l["teardown"], l["notes"]]
    for c, v in enumerate(vals, 1):
        cell = wsL.cell(row=r, column=c, value=v)
        cell.font, cell.alignment, cell.border = BODY_FONT, WRAP, BORDER
    wsL.cell(row=r, column=18, value=f'=IF(Q{r}>=60,"Priority",IF(Q{r}>=40,"Warm","Nurture"))')
    for col in (9, 13):
        cell = wsL.cell(row=r, column=col)
        if cell.value:
            cell.hyperlink = cell.value
            cell.font = LINK_FONT
    wsL.cell(row=r, column=10).fill = INPUT_FILL
    wsL.cell(row=r, column=9).fill = INPUT_FILL
    wsL.row_dimensions[r].height = 96
LAST_L = len(leads) + 1
wsL.freeze_panes = "E2"
wsL.auto_filter.ref = f"A1:{get_column_letter(len(LCOLS))}{LAST_L}"
dvYN = DataValidation(type="list", formula1="=Lists!$C$2:$C$3", allow_blank=True)
wsL.add_data_validation(dvYN)
dvYN.add(f"J2:J{LAST_L + 200}")
wsL.conditional_formatting.add(f"J2:J{LAST_L + 200}", CellIsRule(operator="equal", formula=['"No"'], fill=PatternFill("solid", fgColor="F8D7DA")))
wsL.conditional_formatting.add(f"R2:R{LAST_L + 200}", CellIsRule(operator="equal", formula=['"Priority"'], fill=PatternFill("solid", fgColor="D1FADF")))

# ---------- Tracker ----------
wsT = wb.create_sheet("Tracker", 2)
TCOLS = [("ID", 7), ("Name", 22), ("Company", 22), ("Region", 10), ("Tier", 9), ("Activity pattern", 17),
         ("Profile confirmed", 10), ("Status", 16), ("Warm-up started", 12), ("Connection sent", 12),
         ("Accepted", 12), ("First DM sent", 12), ("FU1 due", 12), ("FU1 sent", 12), ("FU2 due", 12),
         ("FU2 sent", 12), ("Replied", 12), ("Teardown sent", 12), ("Call booked", 12), ("Call held (Y/N)", 10),
         ("Outcome", 11), ("Next action", 30), ("Next action date", 12), ("Days since last touch", 10), ("Notes", 40)]
header(wsT, TCOLS)
lookup = {2: "D", 3: "F", 4: "G", 5: "R", 6: "N", 7: "J"}  # Tracker col -> Leads col
DATE_COLS = [9, 10, 11, 12, 14, 16, 17, 18, 19, 23]
INPUT_COLS = [8] + DATE_COLS + [20, 21, 22, 25]
for r, l in enumerate(leads, 2):
    wsT.cell(row=r, column=1, value=l["id"])
    for c, lc in lookup.items():
        wsT.cell(row=r, column=c, value=f'=IFERROR(INDEX(Leads!${lc}:${lc},MATCH($A{r},Leads!$A:$A,0)),"")')
    wsT.cell(row=r, column=8, value="Nurture" if l["score"] < 40 else "Not started")
    wsT.cell(row=r, column=13, value=f'=IF(L{r}="","",L{r}+4)')
    wsT.cell(row=r, column=15, value=f'=IF(N{r}="","",N{r}+6)')
    wsT.cell(row=r, column=24, value=f'=IF(MAX(I{r}:L{r},N{r},P{r}:S{r})=0,"",TODAY()-MAX(I{r}:L{r},N{r},P{r}:S{r}))')
    for c in range(1, len(TCOLS) + 1):
        cell = wsT.cell(row=r, column=c)
        cell.font, cell.border = BODY_FONT, BORDER
        cell.alignment = WRAP if c in (22, 25) else TOP
        if c in INPUT_COLS:
            cell.fill = INPUT_FILL
        if c in DATE_COLS + [13, 15]:
            cell.number_format = "yyyy-mm-dd"
LAST_T = len(leads) + 1
wsT.freeze_panes = "C2"
wsT.auto_filter.ref = f"A1:{get_column_letter(len(TCOLS))}{LAST_T}"
dvS = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${len(STATUSES) + 1}", allow_blank=True)
dvO = DataValidation(type="list", formula1=f"=Lists!$B$2:$B${len(OUTCOMES)}", allow_blank=True)
dvY = DataValidation(type="list", formula1="=Lists!$C$2:$C$3", allow_blank=True)
dvD = DataValidation(type="date", operator="greaterThan", formula1="45000", allow_blank=True,
                     error="Enter a date like 2026-09-29", errorTitle="Date")
for dv in (dvS, dvO, dvY, dvD):
    wsT.add_data_validation(dv)
dvS.add(f"H2:H{LAST_T + 200}")
dvO.add(f"U2:U{LAST_T + 200}")
dvY.add(f"T2:T{LAST_T + 200}")
for c in DATE_COLS:
    col = get_column_letter(c)
    dvD.add(f"{col}2:{col}{LAST_T + 200}")
red = PatternFill("solid", fgColor="F8D7DA")
wsT.conditional_formatting.add(f"M2:M{LAST_T + 200}", FormulaRule(formula=[f'AND(M2<>"",N2="",M2<TODAY())'], fill=red))
wsT.conditional_formatting.add(f"O2:O{LAST_T + 200}", FormulaRule(formula=[f'AND(O2<>"",P2="",O2<TODAY())'], fill=red))
wsT.conditional_formatting.add(f"H2:H{LAST_T + 200}", FormulaRule(formula=['OR(H2="Call booked",H2="Call held",H2="Won")'], fill=PatternFill("solid", fgColor="D1FADF")))

# ---------- Dashboard ----------
wd = wb.create_sheet("Dashboard", 3)
wd.column_dimensions["A"].width = 24
for col in "BCDEFGH":
    wd.column_dimensions[col].width = 14
wd["A1"] = "Pipeline dashboard"
wd["A1"].font = Font(name=FONT, bold=True, size=14)
wd["A2"] = "Updates from the Tracker tab. Rates divide by leads contacted (status past Not started / Warming up / Nurture)."
wd["A2"].font = Font(name=FONT, size=9, italic=True, color="667085")

T = f"Tracker!$H$2:$H${LAST_T + 200}"
wd["A4"], wd["B4"] = "Status", "Leads"
for c in ("A4", "B4"):
    wd[c].font, wd[c].fill = HEAD_FONT, HEAD_FILL
for i, s in enumerate(STATUSES, 5):
    wd.cell(row=i, column=1, value=s).font = BODY_FONT
    wd.cell(row=i, column=2, value=f'=COUNTIF({T},A{i})').font = BODY_FONT
end_s = 4 + len(STATUSES)
wd.cell(row=end_s + 1, column=1, value="Total").font = Font(name=FONT, bold=True, size=10)
wd.cell(row=end_s + 1, column=2, value=f"=SUM(B5:B{end_s})").font = Font(name=FONT, bold=True, size=10)

# funnel stages as cumulative "reached at least" counts
def reached(stage_list):
    return "+".join(f'COUNTIF({T},"{s}")' for s in stage_list)

later = STATUSES[2:]  # from Connection sent onwards
contacted = [s for s in later if s not in ("Nurture",)]
connected = ["Connected", "DM sent", "Follow-up 1 sent", "Follow-up 2 sent", "Replied", "Teardown sent", "Call booked", "Call held", "Proposal sent", "Won", "Lost"]
replied = ["Replied", "Teardown sent", "Call booked", "Call held", "Proposal sent", "Won", "Lost"]
calls = ["Call booked", "Call held", "Proposal sent", "Won", "Lost"]
won = ["Won"]
r0 = end_s + 4
wd.cell(row=r0 - 1, column=1, value="Funnel").font = Font(name=FONT, bold=True, size=12)
for c, t in enumerate(["Stage", "Leads", "Rate vs contacted"], 1):
    cell = wd.cell(row=r0, column=c, value=t)
    cell.font, cell.fill = HEAD_FONT, HEAD_FILL
funnel = [("Contacted", contacted), ("Connected", connected), ("Replied", replied), ("Call booked", calls), ("Won", won)]
for i, (label, lst) in enumerate(funnel, r0 + 1):
    wd.cell(row=i, column=1, value=label).font = BODY_FONT
    wd.cell(row=i, column=2, value="=" + reached(lst)).font = BODY_FONT
    c = wd.cell(row=i, column=3, value=f"=IFERROR(B{i}/$B${r0 + 1},0)")
    c.number_format, c.font = "0.0%", BODY_FONT

def breakdown(start_row, title, key_col_letter, keys):
    wd.cell(row=start_row - 1, column=1, value=title).font = Font(name=FONT, bold=True, size=12)
    heads = [title.split(" by ")[-1].capitalize(), "Leads", "Contacted", "Replied", "Calls", "Reply rate", "Call rate"]
    for c, t in enumerate(heads, 1):
        cell = wd.cell(row=start_row, column=c, value=t)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
    K = f"Tracker!${key_col_letter}$2:${key_col_letter}${LAST_T + 200}"
    def cnt(lst, r):
        return "+".join(f'COUNTIFS({K},$A{r},{T},"{s}")' for s in lst)
    for i, k in enumerate(keys, start_row + 1):
        wd.cell(row=i, column=1, value=k).font = BODY_FONT
        wd.cell(row=i, column=2, value=f'=COUNTIF({K},$A{i})').font = BODY_FONT
        wd.cell(row=i, column=3, value="=" + cnt(contacted, i)).font = BODY_FONT
        wd.cell(row=i, column=4, value="=" + cnt(replied, i)).font = BODY_FONT
        wd.cell(row=i, column=5, value="=" + cnt(calls, i)).font = BODY_FONT
        for col, num in ((6, "D"), (7, "E")):
            c = wd.cell(row=i, column=col, value=f"=IFERROR({num}{i}/C{i},0)")
            c.number_format, c.font = "0.0%", BODY_FONT
    return start_row + len(keys) + 3

nr = breakdown(r0 + len(funnel) + 4, "Leads by region", "D", ["India", "UK", "Canada", "US"])
patterns = sorted({l["pattern"] for l in leads})
nr = breakdown(nr, "Leads by activity pattern", "F", patterns)
breakdown(nr, "Leads by tier", "E", ["Priority", "Warm", "Nurture"])

# ---------- Sequence ----------
wq = wb.create_sheet("Sequence", 4)
wq.column_dimensions["A"].width = 24
wq.column_dimensions["B"].width = 95
seq = [
    ("Outreach sequence", ""),
    ("", ""),
    ("Day 1 to 5", "Follow them. Leave 1 or 2 real comments on their funding post or the company page's announcement. If there's nothing recent, just view the profile."),
    ("Day 6", "Send the connection note from the Leads tab."),
    ("On accept", "Send the First DM. One question, no pitch."),
    ("+3 to 4 days, no reply", "Follow-up 1: offer the free 3-point profile teardown (text below)."),
    ("+5 to 6 days after that", "Follow-up 2: last nudge that leaves the door open (text below)."),
    ("Yes to teardown", "Send a 2 to 3 minute Loom within 24 hours covering the teardown focus points, then ask for a 20-minute call."),
    ("", ""),
    ("Follow-up 1", "No worries if this is a busy stretch, [Name].\n\nI put together 3 quick notes on your profile.\nThings I'd change so people who hear about [Company] and look you up land on something that sells it.\n\nWant me to send them over? Takes 2 minutes to watch."),
    ("Follow-up 2", "Last one from me, [Name].\n\nIf LinkedIn moves up the list after the raise settles, I'm around.\nEither way, good luck with [the next milestone]."),
    ("", ""),
    ("Send times", "India 9 to 11am IST. UK 8 to 10am UK time. Canada and US East 8 to 10am ET. US West 8 to 10am PT. Tuesday to Thursday."),
    ("Limits", "Keep connection requests around 15 to 25 a day. Send by hand, no automation."),
    ("Before each send", "30 seconds on their recent activity: still in role, last post date, any agency or ghostwriter already visible. Fill any [bracketed] line yourself after reading the post."),
    ("Email rules", "Canada: CASL (relevant to role, unsubscribe). UK: PECR and UK GDPR. US: CAN-SPAM. India: DPDP Act."),
]
for r, (a, bval) in enumerate(seq, 1):
    ca = wq.cell(row=r, column=1, value=a)
    ca.font = Font(name=FONT, bold=True, size=14 if r == 1 else 10)
    ca.alignment = TOP
    cb = wq.cell(row=r, column=2, value=bval)
    cb.font, cb.alignment = BODY_FONT, WRAP
    if "\n" in bval:
        wq.row_dimensions[r].height = 15 * (bval.count("\n") + 1)

wb.move_sheet("Lists", offset=10)
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT, len(leads), "leads")
